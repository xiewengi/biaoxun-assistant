"""采购人-区县基础资料：CRUD + Excel 批量导入。

buyer 字段存储：
- buyer 列：AES-256-GCM 密文（base64），每次加密随机 IV，相同明文密文不同
- buyer_hash 列：SHA-256(明文)，用于唯一约束 + 匹配查找（不可逆）

所有 API 出入参都是明文 buyer，加解密在本路由内部完成。
"""
from __future__ import annotations

import io
from datetime import datetime
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from pydantic import BaseModel

from .. import crypto, db, security

router = APIRouter(prefix="/api/buyer-districts", tags=["buyer-districts"])


class BuyerDistrictBody(BaseModel):
    buyer: str
    district: str = ""
    remark: str = ""


class BuyerDistrictUpdate(BaseModel):
    buyer: Optional[str] = None
    district: Optional[str] = None
    remark: Optional[str] = None


# ---------- 辅助 ----------
def _row_to_api(row: Dict[str, Any]) -> Dict[str, Any]:
    buyer_plain = crypto.decrypt(row.get("buyer") or "") or "(解密失败)"
    return {
        "id": row["id"],
        "buyer": buyer_plain,
        "district": row.get("district") or "",
        "remark": row.get("remark") or "",
        "createdAt": row.get("created_at") or "",
        "updatedAt": row.get("updated_at") or "",
    }


def _lookup_buyer(buyer: str) -> Optional[Dict[str, Any]]:
    """按明文 buyer 查找（用 SHA-256 哈希匹配密文表里的 buyer_hash 列）。"""
    if not buyer:
        return None
    h = crypto.sha256_hex(buyer)
    return db.query_one(
        "SELECT * FROM buyer_district WHERE buyer_hash = ?", (h,)
    )


# ---------- 路由 ----------
@router.get("")
def list_buyer_districts(
    keyword: str = "",
    page: int = 1,
    page_size: int = 50,
    user: Dict[str, Any] = Depends(security.get_current_user),
) -> Dict[str, Any]:
    security.require_permission(user, "buyers:read", "访问区县映射")
    conditions: List[str] = ["1=1"]
    params: List[Any] = []
    if keyword:
        like = f"%{keyword}%"
        # buyer 列是密文，不能直接 LIKE；对明文 buyer 也做不了模糊匹配（哈希不可逆）
        # 所以 keyword 只能匹配 district / remark 明文列
        conditions.append("(district LIKE ? OR remark LIKE ?)")
        params.extend([like, like])

    where = "WHERE " + " AND ".join(conditions)
    total_row = db.query_one(
        f"SELECT COUNT(*) AS c FROM buyer_district {where}", params
    )
    total = int(total_row["c"] or 0)

    page = max(1, int(page or 1))
    page_size = max(0, int(page_size if page_size is not None else 50))

    if page_size == 0:
        rows = db.query_all(
            f"SELECT * FROM buyer_district {where} ORDER BY id DESC", params
        )
        page_size = total
    else:
        start = (page - 1) * page_size
        rows = db.query_all(
            f"SELECT * FROM buyer_district {where} ORDER BY id DESC LIMIT ? OFFSET ?",
            params + [page_size, start],
        )

    return {
        "items": [_row_to_api(r) for r in rows],
        "total": total,
        "page": page,
        "pageSize": page_size,
    }


@router.post("")
def create_buyer_district(
    body: BuyerDistrictBody,
    user: Dict[str, Any] = Depends(security.get_current_user),
) -> Dict[str, Any]:
    security.require_permission(user, "buyers:manage", "管理区县映射")
    buyer = (body.buyer or "").strip()
    if not buyer:
        raise HTTPException(status_code=400, detail="采购人不能为空")

    existing = _lookup_buyer(buyer)
    if existing:
        raise HTTPException(status_code=400, detail=f"采购人「{buyer}」已存在（id={existing['id']}）")

    db.execute(
        "INSERT INTO buyer_district (buyer, buyer_hash, district, remark) VALUES (?, ?, ?, ?)",
        (crypto.encrypt(buyer), crypto.sha256_hex(buyer),
         (body.district or "").strip(), (body.remark or "").strip()),
    )
    security.audit(user, "新增采购人基础资料", buyer, body.district)
    new_row = _lookup_buyer(buyer)
    return _row_to_api(new_row) if new_row else {"buyer": buyer}


@router.put("/{item_id}")
def update_buyer_district(
    item_id: int,
    body: BuyerDistrictUpdate,
    user: Dict[str, Any] = Depends(security.get_current_user),
) -> Dict[str, Any]:
    security.require_permission(user, "buyers:manage", "管理区县映射")
    row = db.query_one("SELECT * FROM buyer_district WHERE id = ?", (item_id,))
    if not row:
        raise HTTPException(status_code=404, detail="记录不存在")

    # 如果改了 buyer 名字，检查新名字是否冲突
    old_buyer = crypto.decrypt(row["buyer"]) or ""
    new_buyer = body.buyer.strip() if body.buyer is not None else old_buyer
    if body.buyer is not None and new_buyer != old_buyer and new_buyer:
        conflict = _lookup_buyer(new_buyer)
        if conflict and conflict["id"] != item_id:
            raise HTTPException(status_code=400, detail=f"采购人「{new_buyer}」已存在")

    updates: List[str] = []
    params: List[Any] = []
    if body.buyer is not None:
        updates.append("buyer = ?")
        updates.append("buyer_hash = ?")
        params.append(crypto.encrypt(new_buyer))
        params.append(crypto.sha256_hex(new_buyer))
    if body.district is not None:
        updates.append("district = ?")
        params.append(body.district.strip())
    if body.remark is not None:
        updates.append("remark = ?")
        params.append(body.remark.strip())

    if updates:
        updates.append("updated_at = datetime('now', 'localtime')")
        params.append(item_id)
        db.execute(
            f"UPDATE buyer_district SET {', '.join(updates)} WHERE id = ?", params
        )
        security.audit(user, "修改采购人基础资料", new_buyer)

    final = db.query_one("SELECT * FROM buyer_district WHERE id = ?", (item_id,))
    return _row_to_api(final) if final else {"id": item_id}


@router.delete("/{item_id}")
def delete_buyer_district(
    item_id: int,
    user: Dict[str, Any] = Depends(security.get_current_user),
) -> Dict[str, Any]:
    security.require_permission(user, "buyers:manage", "管理区县映射")
    row = db.query_one("SELECT buyer FROM buyer_district WHERE id = ?", (item_id,))
    if not row:
        raise HTTPException(status_code=404, detail="记录不存在")
    buyer_plain = crypto.decrypt(row["buyer"]) or "(解密失败)"
    db.execute("DELETE FROM buyer_district WHERE id = ?", (item_id,))
    security.audit(user, "删除采购人基础资料", buyer_plain)
    return {"ok": True, "id": item_id}


# ---------- Excel 导入 ----------

@router.post("/import-excel")
async def import_excel(
    file: UploadFile = File(...),
    skip_duplicates: bool = Form(True),
    user: Dict[str, Any] = Depends(security.get_current_user),
) -> Dict[str, Any]:
    """上传 Excel，批量导入采购人-区县。

    - skip_duplicates=true（默认）：已存在的采购人跳过
    - skip_duplicates=false：已存在的采购人覆盖更新（以 Excel 为准）
    """
    security.require_permission(user, "buyers:manage", "导入采购人基础资料")

    filename = (file.filename or "").lower()
    if not filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="只支持 .xlsx / .xls 格式")

    try:
        from openpyxl import load_workbook
    except ImportError:
        raise HTTPException(status_code=500, detail="服务端缺少 openpyxl 依赖")

    content = await file.read()
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Excel 解析失败：{e}")

    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)

    header = next(rows_iter, None)
    if not header:
        raise HTTPException(status_code=400, detail="Excel 为空，请检查文件")

    header_list = [str(c).strip() if c is not None else "" for c in header]
    header_idx: Dict[str, int] = {}
    for i, h in enumerate(header_list):
        if h in ("采购人", "名称", "buyer"):
            header_idx["buyer"] = i
        elif h in ("区县", "district"):
            header_idx["district"] = i
        elif h in ("备注", "remark"):
            header_idx["remark"] = i

    if "buyer" not in header_idx:
        raise HTTPException(
            status_code=400,
            detail=f"表头缺少「采购人」列，当前表头：{header_list}",
        )

    buyer_col = header_idx["buyer"]
    district_col = header_idx.get("district")
    remark_col = header_idx.get("remark")

    created = 0
    updated = 0
    skipped = 0
    errors: List[Dict[str, Any]] = []

    # 现有记录用 buyer_hash 做字典，跳过加密解密开销
    existing_map: Dict[str, int] = {}  # buyer_hash -> id
    for row in db.query_all("SELECT id, buyer_hash FROM buyer_district"):
        existing_map[row["buyer_hash"]] = row["id"]

    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    with db.get_conn() as conn:
        for line_no, row in enumerate(rows_iter, start=2):
            if not row or all(c is None or str(c).strip() == "" for c in row):
                continue

            buyer_val = row[buyer_col] if buyer_col < len(row) else None
            buyer = str(buyer_val).strip() if buyer_val is not None else ""
            if not buyer:
                errors.append({"line": line_no, "reason": "采购人不能为空"})
                continue

            district_val = row[district_col] if district_col is not None and district_col < len(row) else None
            district = str(district_val).strip() if district_val is not None else ""

            remark_val = row[remark_col] if remark_col is not None and remark_col < len(row) else None
            remark = str(remark_val).strip() if remark_val is not None else ""

            h = crypto.sha256_hex(buyer)
            existing_id = existing_map.get(h)
            if existing_id is None:
                conn.execute(
                    "INSERT INTO buyer_district (buyer, buyer_hash, district, remark) VALUES (?, ?, ?, ?)",
                    (crypto.encrypt(buyer), h, district, remark),
                )
                existing_map[h] = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
                created += 1
            elif skip_duplicates:
                skipped += 1
            else:
                conn.execute(
                    "UPDATE buyer_district SET buyer = ?, buyer_hash = ?, district = ?, remark = ?, updated_at = ? WHERE id = ?",
                    (crypto.encrypt(buyer), h, district, remark, now_str, existing_id),
                )
                updated += 1

    wb.close()
    security.audit(
        user, "导入采购人基础资料", file.filename or "excel",
        f"新增{created}，更新{updated}，跳过{skipped}，异常{len(errors)}"
    )
    return {
        "ok": True,
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "errors": errors,
    }


@router.get("/match-by-buyer")
def match_by_buyer(
    buyer: str,
    user: Dict[str, Any] = Depends(security.get_current_user),
) -> Dict[str, Any]:
    security.require_permission(user, "buyers:read", "访问区县映射")
    """工具接口：给一个明文采购人，返回匹配到的区县（供公告自动补全用）。"""
    row = _lookup_buyer((buyer or "").strip())
    if not row:
        return {"matched": False, "district": ""}
    return {"matched": True, "district": row.get("district") or ""}
