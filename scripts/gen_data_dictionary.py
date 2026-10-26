#!/usr/bin/env python3
"""生成标讯助手数据字典（Markdown + Excel）。

输入：`dump_schema.py` 导出的 JSON（PRAGMA table_info/index_list/foreign_key_list + 行数）。
用法：
    python3 scripts/gen_data_dictionary.py --schema /tmp/bx-schema.json \
        --out-md docs/data-dictionary.md --out-xlsx docs/data-dictionary.xlsx

说明：字段中文说明维护在本文件的 TABLE_META / COLUMN_META 中；表结构变化后重新导出 schema 即可。
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any, Dict, List, Optional

TABLE_META: Dict[str, Dict[str, str]] = {
    "groups": {
        "title": "小组（团队）",
        "desc": "业务小组及管理员；顺序即团队页展示优先级，集团客户部固定为 id=1、不可删除。",
    },
    "group_members": {
        "title": "小组成员",
        "desc": "小组下的业务员名单，编辑小组时整体替换；同一小组内姓名唯一。",
    },
    "users": {
        "title": "登录账号",
        "desc": "由小组管理员派生的登录账号；口令用 PBKDF2-HMAC-SHA256 加盐存储。",
    },
    "sessions": {
        "title": "登录会话",
        "desc": "Bearer Token 会话；登录时清理过期记录，改密会删除该账号全部会话。",
    },
    "announcements": {
        "title": "招标公告",
        "desc": "核心业务表：公告正文摘要、商机字段、当前处理人与进度；采集器按 id 增量入库。",
    },
    "progress_records": {
        "title": "进度记录",
        "desc": "公告每一次进度变更的流水（分配、确认、手工更新），按 id 升序展示。",
    },
    "assignment_records": {
        "title": "任务分配记录",
        "desc": "公告每次分配给小组/接收人的历史流水。",
    },
    "sources": {
        "title": "数据来源",
        "desc": "采集站点注册表 + 最近一次探活结果；采集器读取同一份注册信息。",
    },
    "schedule_config": {
        "title": "定时任务配置",
        "desc": "单行配置表（id 恒为 1）：关键词、采集间隔、提醒间隔、运行状态、采集窗口。",
    },
    "crawl_runs": {
        "title": "采集运行记录",
        "desc": "每次采集（页面按钮 / API / systemd 定时任务）的运行结果与统计。",
    },
    "audit_log": {
        "title": "操作审计",
        "desc": "登录、确认商机、分配、团队/来源/配置变更等写操作留痕。",
    },
}

# 通用字段说明（跨表复用）
COLUMN_META: Dict[str, str] = {
    "id": "主键；公告表为业务编号（ANN+日期+序号）",
    "created_at": "创建时间（本地时间，YYYY-MM-DD HH:MM:SS）",
    "updated_at": "最后更新时间",
    "ann_id": "所属公告编号，关联 announcements.id",
    "group_id": "所属小组，关联 groups.id",
    "user_id": "所属账号，关联 users.id",
    "note": "备注 / 跟进说明",
    "position": "岗位名称",
    "contact": "联系方式",
    "name": "名称",
    "role": "角色，当前均为 admin",
    "password_hash": "口令哈希（PBKDF2-HMAC-SHA256，12 万轮）",
    "salt": "口令随机盐（十六进制）",
    "phone": "手机号 / 联系电话",
    "account": "登录账号（唯一）",
    "admin": "小组管理员姓名",
    "admin_contact": "管理员联系方式",
    "admin_account": "管理员登录账号（与 users.account 对应）",
    "name_pinyin": "成员登录/拼音账号，仅作展示与账号参考",
    "progress": "进度：待分配 / 已分配 / 已确认",
    "follower": "操作人（进度记录里的跟进人）",
    "follow_date": "跟进日期（YYYY-MM-DD）",
    "attachment": "附件（预留字段，当前不使用）",
    "assigner": "分配人",
    "assigner_group": "分配人所属小组",
    "assignee": "接收人",
    "assignee_group": "接收小组",
    "assigned_at": "分配时间",
    "title": "标题",
    "enabled": "是否启用：1 启用 / 0 停用",
    "created_at_": "创建时间",
    "status": "状态",
    "message": "说明 / 错误信息（截断至 500 字）",
    "operator": "触发者（cli / 页面操作人）",
    "started_at": "开始时间",
    "finished_at": "结束时间",
    "source_code": "采集适配器：czggzy / gdgpo / ccgp / generic",
    "domain": "站点域名",
    "url": "站点地址（唯一）",
}

# 表内特有字段说明（优先于通用说明）
TABLE_COLUMN_META: Dict[str, Dict[str, str]] = {
    "groups": {
        "name": "小组名称（唯一），如 集团客户部 / 战客中心 / 市区 / 潮安 / 饶平",
        "region": "所属区域，用于按区域分配公告",
        "created_at": "小组创建日期",
    },
    "group_members": {
        "name": "成员姓名",
    },
    "users": {
        "name": "姓名（用于公告处理人、提醒可见性匹配）",
    },
    "sessions": {
        "token": "会话 Token（随机 URL-safe 字符串，主键）",
        "expires_at": "过期时间（默认登录后 7 天）",
    },
    "announcements": {
        "id": "公告编号（主键），采集器生成，如 ANN202610234478",
        "title": "公告完整标题",
        "project_name": "项目名称（列表主展示字段）",
        "project_code": "项目/采购编号，采集去重优先键",
        "buyer": "采购人名称",
        "agency": "采购代理机构",
        "budget": "预算金额（保留原始字符串，含千分位/元）",
        "fund_source": "资金来源",
        "district": "所属区县：潮安 / 市区 / 湘桥 / 饶平 等",
        "ann_type": "公告类型，如 采购公告 / 采购意向公告 / 招标公告",
        "pub_date": "发布日期（YYYY-MM-DD）",
        "deadline": "报名截止日期",
        "bid_open_date": "开标日期；标讯状态由它派生",
        "summary": "项目概况全文（采集正文，可能较长）",
        "contact": "联系方式原始串（三方联系人拼接）",
        "contacts_json": "结构化联系方式 JSON：{purchaser, agency, project}",
        "source_url": "原始公告链接",
        "region": "所属地市（当前为潮州）",
        "purchase_method": "采购方式，如 公开招标 / 竞争性磋商",
        "matched_keywords": "命中的采集关键词，JSON 数组",
        "business_tags": "业务标签，JSON 数组（平台/专线/机房 等）",
        "is_valid_opportunity": "商机字段：是否有效商机（是/否，空=未填）",
        "has_operator_involvement": "商机字段：是否有运营商参与（是/否）",
        "opportunity_code": "商机字段：商机编码（自由填写）",
        "opportunity_status": "商机字段：商机了解情况（自由填写）",
        "will_bid": "商机字段：是否应标（准备应标/已经应标/弃标）",
        "tree_placement": "商机字段：挂树选择",
        "current_handler": "当前处理人姓名（默认区域小组管理员，可指定成员）",
        "handler_group": "当前处理小组名称（逻辑关联 groups.name，无外键）",
        "progress": "业务进度：待分配 / 已分配 / 已确认",
        "next_follow_up": "下次跟进日期（分配时默认 = 今天 + 3 天）",
        "confirmed_at": "商机信息确认时间；未确认为空",
    },
    "progress_records": {
        "progress": "本次进度值（待分配/已分配/已确认）",
    },
    "assignment_records": {},
    "sources": {
        "name": "网站名称",
        "last_status": "最近探活结果：ok / fail / 空（未检测）",
        "last_http_code": "最近探活 HTTP 状态码",
        "last_latency_ms": "最近探活耗时（毫秒）",
        "last_checked_at": "最近探活时间",
        "last_message": "探活说明，如「正常 864ms」「连接失败: ...」",
        "last_found_links": "最近探活发现链接数",
        "last_matched_links": "最近探活命中关键词链接数",
    },
    "schedule_config": {
        "id": "固定为 1（单行配置）",
        "keyword_list": "采集关键词，JSON 数组（页面可增删）",
        "crawl_interval": "采集间隔展示值，如 每2小时",
        "reminder_interval": "提醒运行间隔展示值，如 每天 09:00, 15:00",
        "crawl_status": "运行状态：运行中 / 已暂停（仅展示）",
        "crawl_days": "每次采集回溯天数，默认 3",
    },
    "crawl_runs": {
        "status": "运行状态：running / ok / error",
        "sites_ok": "本次成功站点数",
        "sites_total": "启用站点总数",
        "new_count": "本次新增公告数",
    },
    "audit_log": {
        "user_name": "操作人姓名（cli 表示命令行）",
        "action": "动作，如 登录/确认商机/分配任务/更新小组",
        "target": "操作对象（公告编号/小组名/账号等）",
        "detail": "动作详情",
    },
}

TYPE_HINT = {
    "INTEGER": "整数",
    "TEXT": "文本",
    "REAL": "浮点",
    "BLOB": "二进制",
}

ENUMS: List[Dict[str, str]] = [
    {"field": "announcements.progress / progress_records.progress", "values": "待分配 → 已分配 → 已确认"},
    {"field": "announcements.ann_type", "values": "采购公告 / 采购意向公告 / 招标公告 / 建设工程招标公告 / 采购需求 等（采集器白名单）"},
    {"field": "announcements.will_bid", "values": "准备应标 / 已经应标 / 弃标（可空）"},
    {"field": "announcements.is_valid_opportunity / has_operator_involvement", "values": "是 / 否（可空）"},
    {"field": "sources.last_status", "values": "ok / fail / 空（未检测）"},
    {"field": "sources.source_code", "values": "czggzy（潮州公共资源）/ gdgpo（广东省政府采购）/ ccgp（中国政府采购）/ generic（通用适配器）"},
    {"field": "crawl_runs.status", "values": "running / ok / error"},
    {"field": "users.role", "values": "admin（当前所有账号均为管理员）"},
    {"field": "schedule_config.crawl_status", "values": "运行中 / 已暂停"},
]

RELATIONS = [
    ("groups", "1 : N", "group_members", "group_members.group_id → groups.id（删除小组级联删除成员）"),
    ("groups", "1 : N", "users", "users.group_id → groups.id（删除小组时账号一并删除）"),
    ("users", "1 : N", "sessions", "sessions.user_id → users.id（级联删除）"),
    ("announcements", "1 : N", "progress_records", "progress_records.ann_id → announcements.id（级联删除）"),
    ("announcements", "1 : N", "assignment_records", "assignment_records.ann_id → announcements.id（级联删除）"),
    ("groups", "逻辑关联", "announcements", "announcements.handler_group 存小组名称（无外键，便于历史留痕）"),
]


def load_schema(path: Path) -> Dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def col_comment(table: str, col: str) -> str:
    return (
        TABLE_COLUMN_META.get(table, {}).get(col)
        or COLUMN_META.get(col)
        or ""
    )


def fmt_default(value: Any) -> str:
    if value is None:
        return "—"
    return str(value).replace("\n", " ")


def render_markdown(schema: Dict[str, Any]) -> str:
    tables = schema["tables"]
    total_rows = sum(t["rows"] for t in tables.values())
    lines: List[str] = []
    lines.append("# 标讯助手 · 数据字典\n")
    lines.append("> 本文件由线上 SQLite 库结构自动生成，字段说明维护在 `scripts/gen_data_dictionary.py`。\n")
    lines.append("## 1. 概览\n")
    lines.append("| 项 | 值 |")
    lines.append("| --- | --- |")
    lines.append("| 后端 | FastAPI（Python 3.12） |")
    lines.append("| 数据库 | SQLite %s，WAL 模式 |" % schema.get("sqlite_version", ""))
    lines.append("| 线上路径 | `/opt/biaoxun/data/biaoxun.db` |")
    lines.append("| 表数量 | %d 张 |" % len(tables))
    lines.append("| 当前总行数 | %d 行 |" % total_rows)
    lines.append("| 时间格式 | 日期 `YYYY-MM-DD`，时间 `YYYY-MM-DD HH:MM:SS`（本地时区） |")
    lines.append("| JSON 字段 | 以 TEXT 存储 JSON 字符串：`contacts_json`、`matched_keywords`、`business_tags`、`keyword_list` |")
    lines.append("")
    lines.append("## 2. 表清单\n")
    lines.append("| 表名 | 中文名 | 字段数 | 当前行数 | 用途 |")
    lines.append("| --- | --- | ---: | ---: | --- |")
    for name in sorted(tables):
        meta = TABLE_META.get(name, {"title": name, "desc": ""})
        t = tables[name]
        lines.append(
            "| `%s` | %s | %d | %d | %s |"
            % (name, meta["title"], len(t["columns"]), t["rows"], meta["desc"])
        )
    lines.append("")
    lines.append("## 3. 表关系\n")
    lines.append("| 主表 | 关系 | 从表 | 说明 |")
    lines.append("| --- | --- | --- | --- |")
    for a, rel, b, desc in RELATIONS:
        lines.append("| `%s` | %s | `%s` | %s |" % (a, rel, b, desc))
    lines.append("")
    lines.append("```text")
    lines.append("groups ──1:N──> group_members")
    lines.append("groups ──1:N──> users ──1:N──> sessions")
    lines.append("announcements ──1:N──> progress_records")
    lines.append("announcements ──1:N──> assignment_records")
    lines.append("announcements.handler_group ──逻辑──> groups.name")
    lines.append("（sources / schedule_config / crawl_runs / audit_log 为独立表）")
    lines.append("```")
    lines.append("")
    lines.append("## 4. 表结构明细\n")

    for name in sorted(tables):
        t = tables[name]
        meta = TABLE_META.get(name, {"title": name, "desc": ""})
        pks = [c["name"] for c in t["columns"] if c.get("pk")]
        lines.append("### %s · `%s`\n" % (meta["title"], name))
        lines.append("%s\n" % meta["desc"])
        lines.append("| 项 | 值 |")
        lines.append("| --- | --- |")
        lines.append("| 主键 | %s |" % ("、".join("`%s`" % p for p in pks) or "—"))
        fks = t.get("foreign_keys") or []
        if fks:
            fk_text = "；".join(
                "`%s` → `%s.%s`" % (f["from"], f["table"], f["to"]) for f in fks
            )
        else:
            fk_text = "—"
        lines.append("| 外键 | %s |" % fk_text)
        idx = t.get("indexes") or []
        idx_text = "；".join(
            "%s(%s)%s"
            % (
                "唯一索引" if i["unique"] else "索引",
                ", ".join(i["columns"]),
                "" if i["origin"] == "c" else "（自动）",
            )
            for i in idx
        )
        lines.append("| 索引 | %s |" % (idx_text or "—"))
        lines.append("| 当前行数 | %d |" % t["rows"])
        lines.append("")
        lines.append("| 字段 | 类型 | 必填 | 默认值 | 说明 |")
        lines.append("| --- | --- | --- | --- | --- |")
        for c in t["columns"]:
            notnull = "PK" if c.get("pk") else ("是" if c.get("notnull") else "否")
            lines.append(
                "| `%s` | %s | %s | %s | %s |"
                % (
                    c["name"],
                    TYPE_HINT.get((c.get("type") or "").upper(), c.get("type") or ""),
                    notnull,
                    fmt_default(c.get("dflt_value")),
                    col_comment(name, c["name"]),
                )
            )
        lines.append("")

    lines.append("## 5. 枚举与取值约定\n")
    lines.append("| 字段 | 取值 |")
    lines.append("| --- | --- |")
    for e in ENUMS:
        lines.append("| %s | %s |" % (e["field"], e["values"]))
    lines.append("")
    lines.append("## 6. 运维说明\n")
    lines.append("- **备份**：`server/deploy/backup.sh`（SQLite 在线备份 API，保留 14 天），线上每日 03:15 自动执行。")
    lines.append("- **迁移**：建表脚本 `server/app/schema.sql` 使用 `CREATE TABLE IF NOT EXISTS`；新增字段需手动 `ALTER TABLE` 或重建库。")
    lines.append("- **字符集**：SQLite UTF-8；中文全文检索用 `LIKE`（当前数据量足够）。")
    lines.append("- **删除策略**：小组→成员/账号、公告→记录 均 `ON DELETE CASCADE`；删除小组会同时删除其管理员账号。")
    lines.append("")
    return "\n".join(lines)


def render_xlsx(schema: Dict[str, Any], path: Path) -> Optional[str]:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return "openpyxl 未安装，跳过 Excel"

    tables = schema["tables"]
    wb = Workbook()
    ws = wb.active
    ws.title = "表清单"
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    head_font = Font(bold=True)

    headers = ["表名", "中文名", "字段数", "当前行数", "用途"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = head_fill
        cell.font = head_font
    for name in sorted(tables):
        meta = TABLE_META.get(name, {"title": name, "desc": ""})
        ws.append([name, meta["title"], len(tables[name]["columns"]), tables[name]["rows"], meta["desc"]])
    for i, width in enumerate([24, 18, 10, 10, 70], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws2 = wb.create_sheet("数据字典")
    headers2 = ["表名", "表中文名", "字段", "类型", "主键", "必填", "默认值", "说明", "外键"]
    ws2.append(headers2)
    for cell in ws2[1]:
        cell.fill = head_fill
        cell.font = head_font
    for name in sorted(tables):
        meta = TABLE_META.get(name, {"title": name, "desc": ""})
        fk_map = {f["from"]: "%s.%s" % (f["table"], f["to"]) for f in tables[name].get("foreign_keys") or []}
        for c in tables[name]["columns"]:
            ws2.append(
                [
                    name,
                    meta["title"],
                    c["name"],
                    TYPE_HINT.get((c.get("type") or "").upper(), c.get("type") or ""),
                    "是" if c.get("pk") else "",
                    "是" if c.get("notnull") else "否",
                    fmt_default(c.get("dflt_value")),
                    col_comment(name, c["name"]),
                    fk_map.get(c["name"], ""),
                ]
            )
    for i, width in enumerate([22, 16, 24, 10, 8, 8, 38, 60, 28], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = width
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions
    for row in ws2.iter_rows(min_row=1):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.auto_filter.ref = ws.dimensions

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return None


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--schema", required=True)
    parser.add_argument("--out-md", required=True)
    parser.add_argument("--out-xlsx", default="")
    args = parser.parse_args()
    schema = load_schema(Path(args.schema))
    md = render_markdown(schema)
    out_md = Path(args.out_md)
    out_md.parent.mkdir(parents=True, exist_ok=True)
    out_md.write_text(md, encoding="utf-8")
    print("✅ 已生成：%s" % out_md)
    if args.out_xlsx:
        err = render_xlsx(schema, Path(args.out_xlsx))
        if err:
            print("⚠️ %s" % err)
        else:
            print("✅ 已生成：%s" % args.out_xlsx)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
